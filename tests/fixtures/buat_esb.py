"""Pembuat file ESB tiruan untuk tes.

Bentuknya meniru export ESB asli (Sales Recapitulation Report & Sales Menu
COGS Report, Okt 2026): metadata di baris 1-11, header di baris 13, data,
lalu khusus Bill Report satu baris footer tanpa label yang berisi total
sebagai teks berformat Indonesia di bawah kolomnya. COGS Report tidak punya
footer. Satu file bisa berisi beberapa cabang (kolom Branch per baris).
"""

from __future__ import annotations

from datetime import date
from decimal import Decimal
from pathlib import Path

import openpyxl

KOLOM_BILL = ["Sales Number", "Bill Number", "Sales Type", "Sales Date", "Sales In Date",
              "Sales In Time", "Sales Out Date", "Sales Out Time", "Branch", "Visit Purpose", "Table",
              "Customer Name", "Promotion", "Pax Total", "Subtotal", "Menu Discount", "Bill Discount",
              "Voucher Discount", "Net Sales", "Service Charge Total", "Tax Total", "VAT Total",
              "Rounding Total", "Grand Total", "DPP", "Cashier"]
KOLOM_COGS = ["Sales Number", "Sales Date", "Sales Type", "Branch", "Menu", "Menu Code",
              "Menu Category", "Menu Category Detail", "Qty", "Price", "Total", "Discount Total",
              "COGS Total", "COGS Total (%)", "Margin"]
FOOTER_BILL = ["Pax Total", "Subtotal", "Menu Discount", "Bill Discount", "Voucher Discount",
               "Net Sales", "Service Charge Total", "Tax Total", "Rounding Total", "Grand Total"]


def rp_id(n) -> str:
    """Decimal -> '231.973.850,00'."""
    d = Decimal(n).quantize(Decimal("0.01"))
    bulat, _, sen = f"{abs(d):f}".partition(".")
    return ("-" if d < 0 else "") + f"{int(bulat):,}".replace(",", ".") + "," + sen


def _tulis(path: Path, judul: str, kolom: list[str], baris: list[dict], footer: dict | None,
           cabang, awal: date, akhir: date, baris_header: int = 13) -> Path:
    wb = openpyxl.Workbook()
    ws = wb.active
    daftar = [cabang] if isinstance(cabang, str) else list(cabang)
    meta = [
        [judul],
        ["KEDAI AMPYANG"],
        [],
        ["Generated", "06-10-2026 11:10:35"],
        ["Period", f"{awal:%d-%m-%Y} - {akhir:%d-%m-%Y}"],
        ["Branch", ", ".join(f"Kedai Ampyang - {c}" for c in daftar)],
        ["Sales Type", "Sales, Non Sales"],
        ["Company", "KEDAI AMPYANG"],
    ]
    for i, r in enumerate(meta, start=1):
        for j, v in enumerate(r, start=1):
            ws.cell(i, j, v)
    for j, k in enumerate(kolom, start=1):
        ws.cell(baris_header, j, k)
    n = baris_header
    for r in baris:
        n += 1
        for j, k in enumerate(kolom, start=1):
            v = r.get(k)
            ws.cell(n, j, float(v) if isinstance(v, Decimal) else v)
    if footer is not None:  # baris footer ESB: tanpa label, angka di bawah kolomnya
        n += 1
        for j, k in enumerate(kolom, start=1):
            if k in footer:
                ws.cell(n, j, rp_id(footer[k]))
    wb.save(path)
    return path


def footer_dari(baris: list[dict], kolom: list[str]) -> dict:
    return {k: sum((Decimal(str(r.get(k) or 0)) for r in baris), Decimal(0)) for k in kolom}


TANPA_FOOTER = object()


def buat_bill(path, baris, cabang="Rungkut", awal=date(2026, 9, 1), akhir=date(2026, 9, 7),
              footer=None, **kw) -> Path:
    if footer is TANPA_FOOTER:
        footer = None
    elif footer is None:
        footer = footer_dari(baris, FOOTER_BILL)
    return _tulis(Path(path), "Sales Recapitulation Report", KOLOM_BILL, baris, footer,
                  cabang, awal, akhir, **kw)


def buat_cogs(path, baris, cabang="Rungkut", awal=date(2026, 9, 1), akhir=date(2026, 9, 7), **kw) -> Path:
    return _tulis(Path(path), "Sales Menu COGS Report", KOLOM_COGS, baris, None, cabang, awal, akhir, **kw)


def bill(no, tgl, jam="12:00:00", sub=50000, pax=1, vp="DINE IN", disk=0, cabang="Rungkut",
         sc=0, tipe="Sales", **kw) -> dict:
    """sc = service charge dalam persen (Rungkut 3%, Mawar 0% di data asli)."""
    sub = Decimal(sub)
    net = sub - Decimal(disk)
    servis = net * Decimal(sc) / 100
    pajak = (net + servis) * Decimal("0.1")
    r = {"Sales Number": no, "Bill Number": f"B{no}", "Sales Type": tipe, "Sales Date": tgl,
         "Sales In Date": tgl, "Sales In Time": jam, "Sales Out Date": tgl, "Sales Out Time": jam,
         "Branch": f"Kedai Ampyang - {cabang}", "Visit Purpose": vp, "Table": "1",
         "Customer Name": "-", "Promotion": None, "Pax Total": pax, "Subtotal": sub,
         "Menu Discount": Decimal(disk), "Bill Discount": Decimal(0), "Voucher Discount": Decimal(0),
         "Net Sales": net, "Service Charge Total": servis, "Tax Total": pajak, "VAT Total": Decimal(0),
         "Rounding Total": Decimal(0), "Grand Total": net + servis + pajak, "DPP": Decimal(0),
         "Cashier": "Kasir A"}
    r.update(kw)
    return r


def item(no, tgl, menu, qty, price, cogs, kat="FOOD", detail="BUBUR & NASI", cabang="Rungkut",
         tipe="Sales", **kw) -> dict:
    total = Decimal(qty) * Decimal(price)
    r = {"Sales Number": no, "Sales Date": tgl, "Sales Type": tipe, "Branch": f"Kedai Ampyang - {cabang}",
         "Menu": menu, "Menu Code": None,
         "Menu Category": kat, "Menu Category Detail": detail, "Qty": qty, "Price": Decimal(price),
         "Total": total, "Discount Total": Decimal(0), "COGS Total": Decimal(cogs),
         "COGS Total (%)": None, "Margin": total - Decimal(cogs)}
    r.update(kw)
    return r
