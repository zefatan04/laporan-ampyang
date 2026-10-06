from datetime import date
from decimal import Decimal

import openpyxl

from app.parser.esb import (BILL, COGS, baca_file_esb, bersihkan_nama_menu, daftar_cabang_metadata,
                            nama_cabang)
from tests.fixtures.buat_esb import TANPA_FOOTER, bill, buat_bill, buat_cogs, item

T1, T2 = date(2026, 9, 1), date(2026, 9, 2)


def test_bill_terbaca(tmp_path):
    p = buat_bill(tmp_path / "b.xlsx", [bill("S1", T1, sub=50000), bill("S2", T2, sub=30000, pax=3, sc=3)])
    h = baca_file_esb(p)
    assert h.jenis is BILL
    assert h.galat == []
    assert h.baris_header == 13
    assert h.periode == (date(2026, 9, 1), date(2026, 9, 7))
    assert daftar_cabang_metadata(h.cabang) == ["Rungkut"]
    assert list(h.data["sales_number"]) == ["S1", "S2"]
    assert h.data["subtotal"].tolist() == [Decimal("50000"), Decimal("30000")]
    assert h.data["sales_date"].tolist() == [T1, T2]
    assert h.data["sales_in_time"].tolist() == ["12:00", "12:00"]
    assert h.data["cabang"].tolist() == ["Rungkut", "Rungkut"]
    # 30.000 + servis 3% (900) + pajak 10% (3.090) = 33.990; 50.000 + 5.000 = 55.000
    assert h.footer["Grand Total"] == Decimal("88990.00")
    assert h.footer["Service Charge Total"] == Decimal("900.00")
    assert h.footer["Pax Total"] == Decimal(4)
    assert h.baris_footer == 16


def test_kolom_nama_pelanggan_tidak_dibaca(tmp_path):
    p = buat_bill(tmp_path / "b.xlsx", [bill("S1", T1, **{"Customer Name": "Bu Ani"})])
    h = baca_file_esb(p)
    assert "Bu Ani" not in h.data.astype(str).to_string()


def test_satu_file_dua_cabang(tmp_path):
    p = buat_bill(tmp_path / "b.xlsx", [bill("S1", T1, cabang="Mawar"), bill("S2", T1, cabang="Rungkut")],
                  cabang=["Mawar", "Rungkut"])
    h = baca_file_esb(p)
    assert h.galat == []
    assert daftar_cabang_metadata(h.cabang) == ["Mawar", "Rungkut"]
    assert h.data["cabang"].tolist() == ["Mawar", "Rungkut"]


def test_cabang_tak_dikenal_jadi_galat(tmp_path):
    p = buat_bill(tmp_path / "b.xlsx", [bill("S1", T1, Branch="Kedai Ampyang - Darmo")])
    assert any("Darmo" in c.pesan for c in baca_file_esb(p).galat)


def test_nama_cabang():
    assert nama_cabang("Kedai Ampyang - Rungkut") == "Rungkut"
    assert nama_cabang("Kedai Ampyang - Mawar, Kedai Ampyang - Rungkut") is None
    assert daftar_cabang_metadata("Kedai Ampyang - Mawar, Kedai Ampyang - Rungkut") == ["Mawar", "Rungkut"]
    assert daftar_cabang_metadata("Kedai X") is None


def test_header_pindah_baris_jadi_peringatan(tmp_path):
    p = buat_bill(tmp_path / "b.xlsx", [bill("S1", T1)], baris_header=11)
    h = baca_file_esb(p)
    assert h.galat == []
    assert h.baris_header == 11
    assert any("baris 11" in c.pesan for c in h.peringatan)


def test_cogs_tanpa_footer_dan_teks_dibersihkan(tmp_path):
    p = buat_cogs(tmp_path / "c.xlsx", [item("S1", T1, "Nasi  Goreng  Jawa ", 1, 25000, 9000,
                                             kat="FOOD ", detail="TEH ")])
    h = baca_file_esb(p)
    assert h.jenis is COGS
    assert h.galat == [] and h.baris_footer is None
    assert h.data["menu_bersih"].tolist() == ["Nasi Goreng Jawa"]
    assert h.data["menu_category"].tolist() == ["FOOD"]
    assert h.data["menu_category_detail"].tolist() == ["TEH"]
    assert bersihkan_nama_menu("  Es   Teh ") == "Es Teh"


def test_uang_berpecahan_tidak_dibulatkan_saat_dibaca(tmp_path):
    # COGS asli: 4183.296, 2583.3333 - dibulatkan per baris akan menggeser total.
    p = buat_cogs(tmp_path / "c.xlsx", [item("S1", T1, "Teh", 1, 0, Decimal("2583.3333"))])
    assert baca_file_esb(p).data["cogs_total"].tolist() == [Decimal("2583.333300")]


def test_jenis_dikenali_dari_isi_bukan_nama(tmp_path):
    p = buat_cogs(tmp_path / "bill-report.xlsx", [item("S1", T1, "Kopi", 1, 10000, 3000)])
    assert baca_file_esb(p).jenis is COGS


def test_file_asing_ditolak(tmp_path):
    p = tmp_path / "x.xlsx"
    wb = openpyxl.Workbook()
    wb.active.append(["Nama", "Alamat"])
    wb.save(p)
    h = baca_file_esb(p)
    assert h.jenis is None and h.galat


def _ubah_sel(p, baris, nama_kolom, nilai):
    wb = openpyxl.load_workbook(p)
    ws = wb.active
    j = [c.value for c in ws[13]].index(nama_kolom) + 1
    ws.cell(baris, j, nilai)
    wb.save(p)


def test_kolom_wajib_hilang(tmp_path):
    p = buat_bill(tmp_path / "b.xlsx", [bill("S1", T1)])
    _ubah_sel(p, 13, "Subtotal", "Sub Total")
    h = baca_file_esb(p)
    assert h.kolom_hilang == ["Subtotal"]
    assert not h.bisa_dipakai


def test_baris_tanpa_id_di_tengah_data_dilaporkan(tmp_path):
    p = buat_bill(tmp_path / "b.xlsx", [bill("S1", T1), bill(None, T1), bill("S3", T1)])
    h = baca_file_esb(p)
    assert any(c.baris == 15 and "tanpa 'Sales Number'" in c.pesan for c in h.galat)


def test_angka_rusak_dilaporkan_bukan_dinolkan(tmp_path):
    p = buat_bill(tmp_path / "b.xlsx", [bill("S1", T1)])
    _ubah_sel(p, 14, "Subtotal", "50,000.00")  # format Amerika
    h = baca_file_esb(p)
    assert any(c.baris == 14 and "Subtotal" in c.pesan for c in h.galat)


def test_bill_tanpa_footer_jadi_galat(tmp_path):
    p = buat_bill(tmp_path / "b.xlsx", [bill("S1", T1)], footer=TANPA_FOOTER)
    assert any("footer" in c.pesan.lower() for c in baca_file_esb(p).galat)
